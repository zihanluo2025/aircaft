"""Build the shift report as a downloadable .xlsx workbook."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from ..schemas import ReportOut

SHIFT_NAME = {"morning": "Morning", "night": "Night"}


def build_report_workbook(report: ReportOut) -> BytesIO:
    bold = Font(bold=True)
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    for label, value in [
        ("Date", report.date.isoformat()),
        ("Shift", SHIFT_NAME.get(report.shift, report.shift)),
        ("Flights done", f"{report.done} / {report.total}"),
        ("Average minutes", report.average_minutes),
        ("Signed", report.signed),
        ("Called in", ", ".join(report.called_in) or "none"),
        ("Roster changes", ", ".join(report.roster_changes) or "none"),
    ]:
        summary_ws.append([label, value])
        summary_ws.cell(row=summary_ws.max_row, column=1).font = bold
    summary_ws.column_dimensions["A"].width = 18
    summary_ws.column_dimensions["B"].width = 70

    flights_ws = wb.create_sheet("Flights")
    flights_ws.append(["Flight", "ETA", "Status", "Minutes", "Staff", "Signed", "Crew"])
    for cell in flights_ws[1]:
        cell.font = bold
    for f in report.flights:
        crew = "; ".join(f"{g.role}: {', '.join(g.names)}" for g in f.roles)
        flights_ws.append([f.flight_no, f.eta, f.status, f.minutes, f.crew_count, "Yes" if f.signed else "No", crew])
    flights_ws.column_dimensions["A"].width = 12
    flights_ws.column_dimensions["C"].width = 10
    flights_ws.column_dimensions["G"].width = 70

    people_ws = wb.create_sheet("Per person")
    people_ws.append(["Name", "Aircraft", "Minutes"])
    for cell in people_ws[1]:
        cell.font = bold
    for p in report.people:
        people_ws.append([p.name, p.aircraft, p.minutes])
    people_ws.column_dimensions["A"].width = 30

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
